import csv
from io import BytesIO, StringIO

from openpyxl import load_workbook

from engines.complexity_engine.complexity import calculate_complexity
from engines.optimizer.prompt_optimizer import optimize_prompt
from engines.token_engine.model_registry import ModelRegistry
from engines.token_engine.token_counter import calculate_cost


MAX_PROMPTS = 500
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
PROMPT_HEADERS = {"prompt", "prompts", "prompt_text", "text", "input"}
COMPLEXITY_METRICS = (
    "word_count",
    "sentence_count",
    "instruction_density",
    "keyword_complexity",
    "entropy",
    "reasoning_indicators",
)


def parse_prompt_file(contents, filename):
    if len(contents) > MAX_UPLOAD_BYTES:
        raise ValueError("The uploaded file exceeds the 5 MB limit.")

    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if extension == "csv":
        try:
            rows = list(csv.reader(StringIO(contents.decode("utf-8-sig"))))
        except UnicodeDecodeError as error:
            raise ValueError("CSV files must use UTF-8 encoding.") from error
        except csv.Error as error:
            raise ValueError("Could not read the CSV file. Check its formatting.") from error
    elif extension == "xlsx":
        try:
            workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows) > MAX_PROMPTS + 1:
                        raise ValueError(f"The spreadsheet may contain at most {MAX_PROMPTS} prompts.")
            finally:
                workbook.close()
        except ValueError:
            raise
        except Exception as error:
            raise ValueError("Could not read the Excel file. Please upload a valid .xlsx workbook.") from error
    else:
        raise ValueError("Upload an .xlsx Excel workbook or a .csv file.")

    if not rows:
        raise ValueError("The uploaded file contains no prompts.")

    first_row = rows[0]
    prompt_column = next(
        (
            index
            for index, value in enumerate(first_row)
            if str(value or "").strip().lower() in PROMPT_HEADERS
        ),
        0,
    )
    has_header = any(
        str(value or "").strip().lower() in PROMPT_HEADERS for value in first_row
    )
    data_rows = rows[1:] if has_header else rows
    prompts = [
        str(row[prompt_column]).strip()
        for row in data_rows
        if len(row) > prompt_column and row[prompt_column] is not None
        and str(row[prompt_column]).strip()
    ]

    if len(prompts) > MAX_PROMPTS:
        raise ValueError(f"The spreadsheet may contain at most {MAX_PROMPTS} prompts.")
    if not prompts:
        raise ValueError("No prompts found. Add a 'prompt' column or put prompts in the first column.")
    return prompts


def analyze_prompts(prompts, model="GPT-4.1"):
    if model not in ModelRegistry().list_models():
        raise ValueError("Select a supported model.")
    if not prompts:
        raise ValueError("Provide at least one prompt.")
    if len(prompts) > MAX_PROMPTS:
        raise ValueError(f"Analyze at most {MAX_PROMPTS} prompts at a time.")

    results = []
    for index, prompt in enumerate(prompts, start=1):
        complexity = calculate_complexity(prompt)
        optimization = optimize_prompt(prompt, model)
        original_tokens = optimization["original_tokens"]
        optimized_tokens = optimization["new_tokens"]
        original_output_tokens = max(20, original_tokens // 3)
        optimized_output_tokens = max(20, optimized_tokens // 3)
        original_cost = calculate_cost(original_tokens, original_output_tokens, model)
        optimized_cost = calculate_cost(optimized_tokens, optimized_output_tokens, model)

        results.append({
            "prompt_number": index,
            "prompt": prompt,
            "optimized_prompt": optimization["optimized_prompt"],
            "complexity_score": complexity["score"],
            "complexity_category": complexity["category"],
            **{metric: complexity[metric] for metric in COMPLEXITY_METRICS},
            "accuracy_evaluation": "Not evaluated: no model outputs or reference answers supplied.",
            "original_tokens": original_tokens,
            "optimized_tokens": optimized_tokens,
            "tokens_saved": original_tokens - optimized_tokens,
            "percentage_saved": optimization["percentage_saved"],
            "estimated_cost_before": original_cost,
            "estimated_cost_after": optimized_cost,
            "estimated_cost_savings": round(original_cost - optimized_cost, 6),
        })

    total_original_tokens = sum(row["original_tokens"] for row in results)
    total_optimized_tokens = sum(row["optimized_tokens"] for row in results)
    total_cost_before = round(sum(row["estimated_cost_before"] for row in results), 6)
    total_cost_after = round(sum(row["estimated_cost_after"] for row in results), 6)
    prompt_count = len(results)
    return {
        "model": model,
        "summary": {
            "prompt_count": prompt_count,
            "average_complexity_score": round(
                sum(row["complexity_score"] for row in results) / prompt_count, 2
            ),
            "high_complexity_prompts": sum(
                row["complexity_category"] == "High" for row in results
            ),
            "prompts_reduced": sum(row["tokens_saved"] > 0 for row in results),
            "original_tokens": total_original_tokens,
            "optimized_tokens": total_optimized_tokens,
            "tokens_saved": total_original_tokens - total_optimized_tokens,
            "percentage_saved": round(
                (total_original_tokens - total_optimized_tokens)
                / total_original_tokens * 100,
                2,
            ) if total_original_tokens else 0.0,
            "estimated_cost_before": total_cost_before,
            "estimated_cost_after": total_cost_after,
            "estimated_cost_savings": round(total_cost_before - total_cost_after, 6),
            "accuracy": "Not evaluated: no model outputs or reference answers supplied.",
        },
        "results": results,
    }


def batch_results_csv(report):
    output = StringIO()
    fieldnames = list(report["results"][0])
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for row in report["results"]:
        writer.writerow({
            key: "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
            for key, value in row.items()
        })
    return output.getvalue()
